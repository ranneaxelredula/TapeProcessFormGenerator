"""
Tape Process Form Generator  –  IBAD-PF-009 Rev 00
----------------------------------------------------
Reads an Excel file (Monday.com / IBAD database export) and generates one
multi-page Word file by filling the Tape Process Form Word template for each
tape row; the result stays editable in Word.

Pages per tape (Monday and RR124 formats alike) depend on the tape's Length
column — "Length, m" in Monday exports, "Length" in RR124 exports:
  - under 600 m        → 1 page
  - 600 m to 1099 m    → 2 pages
  - 1100 m and above   → 3 pages
The first page carries every extracted value including the coordinates; any
additional page repeats the same values but keeps the form's coordinate blank
for hand entry.

Fields filled per page (a missing value leaves the blank for hand entry). The
first five are centred on their blank's line, underlined like the form:
  - Hastelloy thickness   → blank before "microns"
  - Name (#…(low-high))  → tape number + coordinates (each number on its part
                            of "(___ ___-  ____ )")
  - max(Evaluation_start, Evaluation_end) → IBAD#3_EVAL, before the form's "A"
  - Reference Ic          → IBAD#3_ref, before the form's "A"
  - XRD Tilt MIN / AVE    → before the form's "°" (RR124 exports only)
  - IBAD deposition date  → the "202_/ 月/ 日" date of the IBAD Deposition row
  - "Sent to PLD by:"     → that row's Responsible Person column

Usage:
  1. Place this file in the same folder as IBAD-PF-009_Tape Process Form_Rev.0.docx
  2. Double-click TapeProcessFormGenerator.pyw
  3. Browse to your Excel file, confirm the output path, click Generate.

To make a stand-alone TapeProcessFormGenerator.exe (no Python needed to run
it), double-click build_exe.bat; then keep the Word template next to the .exe.
"""

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import threading
import copy
import io
import os
import re
import sys
import unicodedata
from collections import defaultdict

# ── Auto-install missing dependencies ─────────────────────────────────────────
try:
    import pandas as pd
    import docx                                   # package "python-docx"
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    from docx.oxml import OxmlElement
except ImportError:
    # In a built .exe, sys.executable is the .exe itself, so "pip install"
    # would just relaunch the program; a missing module there is a build bug.
    if getattr(sys, "frozen", False):
        raise
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install",
                           "pandas", "openpyxl", "python-docx",
                           "--quiet"])
    import pandas as pd
    import docx
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    from docx.oxml import OxmlElement

# ── Program files ─────────────────────────────────────────────────────────────
TEMPLATE_FILE = "IBAD-PF-009_Tape Process Form_Rev.0.docx"
ICON_FILE     = "logo.ico"     # window icon; build_exe.bat also uses it for the .exe


def _app_dir():
    """Folder the program lives in: the .exe's folder when built, else this script's."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


def _bundled_file(name):
    """Path of a file packed into the .exe (PyInstaller unpacks it to
    sys._MEIPASS), or of the file next to this script when not built."""
    return os.path.join(getattr(sys, "_MEIPASS", None) or _app_dir(), name)


# Input validation limits for the "Sent to PLD by:" field
PLD_MIN_CHARS = 4
PLD_MAX_CHARS = 10

# ── Pages per tape, by Length (metres) ────────────────────────────────────────
# under PAGE_LEN_2          → 1 page
# PAGE_LEN_2 .. PAGE_LEN_3  → 2 pages   (600 m itself gives 2)
# PAGE_LEN_3 and above      → 3 pages   (1100 m itself gives 3)
PAGE_LEN_2 = 600.0
PAGE_LEN_3 = 1100.0
MAX_PAGES  = 3


def pages_for_length(length):
    """
    Number of form pages to generate for a tape of `length` metres.

    A missing / blank / non-numeric length falls back to a single page, which
    is the safe default: the operator gets one form rather than none.
    """
    if length is None:
        return 1
    try:
        L = float(length)
    except (TypeError, ValueError):
        return 1
    if L != L:                       # NaN
        return 1
    if L >= PAGE_LEN_3:
        return 3
    if L >= PAGE_LEN_2:
        return 2
    return 1


# ──────────────────────────────────────────────────────────────────────────────
#  PARSING HELPERS
# ──────────────────────────────────────────────────────────────────────────────

def parse_name(raw_name: str):
    """
    Extract tape identifier and sorted coordinate range from the Name column.

    Examples
    --------
    "#3244(1235-2165) NEW ULBRICH"  →  "3244",  1235, 2165
    "#KM3268(450-10)"               →  "KM3268",  10,  450  (sorted low→high)
    "#KM3264 (25-814)"              →  "KM3264",  25,  814
    Returns (None, None, None) when the pattern is absent.
    """
    m = re.search(r"#([A-Za-z0-9]+)\s*\(\s*(\d+)\s*[-\u2013]\s*(\d+)\s*\)",
                  str(raw_name))
    if not m:
        return None, None, None
    c1, c2 = int(m.group(2)), int(m.group(3))
    return m.group(1), min(c1, c2), max(c1, c2)


def fmt_num(val) -> str:
    """Convert Excel numeric cells to clean strings (drop trailing .0)."""
    try:
        f = float(val)
        if f != f:          # NaN check
            return ""
        return str(int(f)) if f == int(f) else str(round(f, 3))
    except (ValueError, TypeError):
        s = str(val).strip()
        return "" if s.lower() in ("nan", "none", "") else s


def pick_eval(eval_start_raw, eval_end_raw) -> str:
    """
    Return the highest of Evaluation_start and Evaluation_end.
    If only one is available, use that one.
    If both are missing, return empty string.
    """
    s = fmt_num(eval_start_raw)
    e = fmt_num(eval_end_raw)
    if s and e:
        try:
            return s if float(s) >= float(e) else e
        except ValueError:
            return s        # not numeric – fall back to start
    return s or e           # whichever is non-empty, or "" if both empty


def fmt_deposition_date(raw_value):
    """
    Convert the 'IBAD deposition date' cell into (year, month, day) strings.

    The cell is typically a date/timestamp (e.g. 2026-05-17) but may also be a
    plain YYYY-MM-DD string. Month and day are zero-padded to two digits to
    match the form layout, e.g. ("2026", "05", "17").

    Returns (None, None, None) when the cell is empty / not a valid date, so the
    form is left blank for that row.
    """
    if raw_value is None:
        return None, None, None

    # pandas Timestamp / datetime
    ts = pd.to_datetime(raw_value, errors="coerce")
    if pd.isna(ts):
        return None, None, None
    return f"{ts.year:04d}", f"{ts.month:02d}", f"{ts.day:02d}"


def find_header_row(excel_path: str) -> int:
    """Locate the real header row in Monday.com exports (scans first 10 rows)."""
    raw = pd.read_excel(excel_path, engine="openpyxl", header=None, nrows=10)
    for idx, row in raw.iterrows():
        if any(str(v).strip().lower() == "name" for v in row):
            return idx
    return 0


# ──────────────────────────────────────────────────────────────────────────────
#  FORMAT DETECTION + RR124 HELPERS
# ──────────────────────────────────────────────────────────────────────────────

# Signature columns used to recognise an RR124-format export (lower-cased).
_RR124_SIGNATURE = {
    "run_id", "tape_name", "tape_status", "cut_date",
    "sample_cuts_start_coordinate", "sample_cuts_end_coordinate",
}


def detect_format(excel_path: str) -> str:
    """
    Inspect the header row and return 'rr124' or 'monday'.

    Content-based: scans the first 15 rows for either the RR124 signature
    columns or the Monday columns. Raises ValueError if it matches neither.
    """
    raw = pd.read_excel(excel_path, engine="openpyxl", header=None, nrows=15)
    for _, row in raw.iterrows():
        vals = {str(v).strip().lower() for v in row}
        if _RR124_SIGNATURE.issubset(vals):
            return "rr124"
        if "name" in vals and (
            "evaluation_start" in vals or "evaluation_end" in vals
            or "reference ic" in vals
        ):
            return "monday"
    raise ValueError(
        "Unrecognized Excel format.\n\n"
        "The header row matches neither the RR124 columns "
        "(Run_ID, Tape_name, Tape_Status, Cut_Date, Sample_Cuts_*_Coordinate) "
        "nor the Monday columns (Name, Evaluation_start, Reference Ic)."
    )


def _find_rr124_header(excel_path: str) -> int:
    """Locate the RR124 header row (scans first 15 rows for the key columns)."""
    raw = pd.read_excel(excel_path, engine="openpyxl", header=None, nrows=15)
    for idx, row in raw.iterrows():
        vals = {str(v).strip().lower() for v in row}
        if {"run_id", "tape_name", "tape_status"}.issubset(vals):
            return idx
    return 0


def _col(df, name: str):
    """Case-insensitive, whitespace-stripped column lookup."""
    for c in df.columns:
        if str(c).strip().lower() == name.lower():
            return c
    return None


# Accepted headers for the tape-length column (matched like _col): Monday.com
# exports call it "Length, m", RR124 exports call it "Length".
_LENGTH_HEADERS = ("Length", "Length, m")


def _length_col(df):
    """The tape-length column under either export's header, or None."""
    for name in _LENGTH_HEADERS:
        c = _col(df, name)
        if c is not None:
            return c
    return None


# A length typed as text together with its unit, e.g. "855 m", "855m", "650,5 M".
# Group 1 is the number. "mm" (or any other unit) deliberately doesn't match.
_LENGTH_WITH_UNIT = re.compile(r"\s*([\d.,]+)\s*m\s*", re.IGNORECASE)


def parse_length(val):
    """
    Parse a length cell to metres, or None. Same as rr_to_float, except a text
    cell may also carry a trailing "m" unit: "855 m" → 855.0. Full-width
    characters typed with a Japanese IME ("８５５ｍ") are normalized first.
    """
    if isinstance(val, str):
        val = unicodedata.normalize("NFKC", val)
        m = _LENGTH_WITH_UNIT.fullmatch(val)
        if m:
            val = m.group(1)
    return rr_to_float(val)


def _length_and_pages(rec, len_col):
    """
    (length, n_pages) for one row. No length column, or a blank / unreadable
    length, gives (None, 1) — see parse_length and pages_for_length.
    """
    length = parse_length(rec[len_col]) if len_col is not None else None
    return length, pages_for_length(length)


def _status_norm(v) -> str:
    """Normalize Tape_Status: strip + case-fold (so 'cut Eval Sample' counts)."""
    return str(v).strip().casefold()


def _base_name(raw) -> str:
    """Characters before the parenthesis in a Tape_name, whitespace stripped."""
    return str(raw).split("(")[0].strip()


def rr_to_float(val):
    """
    Parse a possibly comma-decimal cell to float (Russian-locale export).
    Returns None for blanks / non-numeric. Only the XRD columns actually use
    comma decimals, but applying this everywhere is harmless.
    """
    if val is None:
        return None
    if isinstance(val, (int, float)):
        f = float(val)
        return None if f != f else f          # NaN → None
    s = str(val).strip()
    if s == "" or s.lower() in ("nan", "none"):
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def fmt_float(v, decimals: int = 2) -> str:
    """Format a float for the form: drop trailing .0, round to `decimals`."""
    if v is None:
        return ""
    if float(v).is_integer():
        return str(int(v))
    return f"{round(float(v), decimals):g}"


def rr124_cut_dates(excel_path: str):
    """Return the sorted distinct Cut_Date values (date objects) on PLD rows."""
    hdr = _find_rr124_header(excel_path)
    df = pd.read_excel(excel_path, engine="openpyxl", header=hdr)
    status_col = _col(df, "Tape_Status")
    cut_col    = _col(df, "Cut_Date")
    if status_col is None or cut_col is None:
        return []
    dates = set()
    for _, row in df.iterrows():
        if _status_norm(row[status_col]) == "sent to pld":
            cd = pd.to_datetime(row[cut_col], errors="coerce")
            if not pd.isna(cd):
                dates.add(cd.date())
    return sorted(dates)


def _cut_endpoints(rec, s_col, e_col):
    """The two coordinate endpoints of a cut-eval row as floats (or None)."""
    return rr_to_float(rec[s_col]), rr_to_float(rec[e_col])


def _dist_to(rec, coord, s_col, e_col) -> float:
    """Distance from `coord` to the nearest endpoint of a cut-eval row."""
    cs, ce = _cut_endpoints(rec, s_col, e_col)
    ds = abs(cs - coord) if cs is not None else float("inf")
    de = abs(ce - coord) if ce is not None else float("inf")
    return min(ds, de)


def _nearest(pool, coord, s_col, e_col, ic_col):
    """
    The cut-eval row nearest to `coord`. Ties broken by higher Ic, then by
    lower coordinate (purely for reproducibility).
    """
    if coord is None or not pool:
        return None

    def key(rec):
        d = _dist_to(rec, coord, s_col, e_col)
        ic = rr_to_float(rec[ic_col])
        ic = ic if ic is not None else 0.0
        cs, ce = _cut_endpoints(rec, s_col, e_col)
        low = min([x for x in (cs, ce) if x is not None], default=float("inf"))
        return (d, -ic, low)        # ascending: nearest, then highest Ic, then lowest coord

    return min(pool, key=key)


def _select_candidates(pool, pld_s, pld_e, s_col, e_col, ic_col):
    """Up to two distinct candidates: nearest to PLD start and to PLD end."""
    picks = []
    for coord in (pld_s, pld_e):
        c = _nearest(pool, coord, s_col, e_col, ic_col)
        if c is not None and not any(c is p for p in picks):
            picks.append(c)
    return picks


def _xrd_triplet(rec, xrd_cols):
    """Return the three XRD tilt floats if all present, else None (all-or-none)."""
    vals = [rr_to_float(rec[c]) for c in xrd_cols]
    if len(vals) >= 3 and all(v is not None for v in vals[:3]):
        return vals[:3]
    return None


def _derive_values(cands, pool, pld_s, pld_e, s_col, e_col, ic_col, ref_col, xrd_cols):
    """
    From the candidate list, derive (eval_val, ref_ic, xrd_min, xrd_ave) strings.

    Selection is NEAREST-first: candidates are ordered by distance to the PLD
    endpoints (ties broken by higher Ic, then lower coordinate). Each field is
    then taken from the nearest candidate that has it, falling back to the next
    nearest. Zero Ic / Reference_Ic is treated as no value (rendered blank).
    """
    if not cands:
        return "", "", "", ""

    def ic_of(rec):
        return rr_to_float(rec[ic_col])

    def cand_dist(rec):
        ds = _dist_to(rec, pld_s, s_col, e_col) if pld_s is not None else float("inf")
        de = _dist_to(rec, pld_e, s_col, e_col) if pld_e is not None else float("inf")
        return min(ds, de)

    def low_coord(rec):
        cs, ce = _cut_endpoints(rec, s_col, e_col)
        return min([x for x in (cs, ce) if x is not None], default=float("inf"))

    # Nearest first; distance tie → higher Ic; then lower coordinate (reproducible).
    ordered = sorted(
        cands,
        key=lambda r: (cand_dist(r),
                       -(ic_of(r) if ic_of(r) is not None else -1.0),
                       low_coord(r)),
    )

    # IC: nearest candidate with a positive Ic (0 / blank treated as no value).
    eval_val = ""
    for r in ordered:
        v = ic_of(r)
        if v is not None and v > 0:
            eval_val = fmt_float(v)
            break

    # XRD triplet: nearest candidate that carries one.
    xrd_min = xrd_ave = ""
    for r in ordered:
        trip = _xrd_triplet(r, xrd_cols)
        if trip:
            xrd_min = fmt_float(min(trip))
            xrd_ave = fmt_float(sum(trip) / len(trip))
            break

    # Reference_Ic: nearest candidate with a positive value; else nearest in pool.
    def ref_of(rec):
        v = rr_to_float(rec[ref_col])
        return v if (v is not None and v > 0) else None

    ref = None
    for r in ordered:
        if ref_of(r) is not None:
            ref = ref_of(r)
            break
    if ref is None:
        with_ref = [r for r in pool if ref_of(r) is not None]
        anchor = pld_s if pld_s is not None else pld_e
        if with_ref and anchor is not None:
            best = min(with_ref, key=lambda r: _dist_to(r, anchor, s_col, e_col))
            ref = ref_of(best)
        elif with_ref:
            ref = ref_of(with_ref[0])
    ref_ic = fmt_float(ref) if ref is not None else ""

    return eval_val, ref_ic, xrd_min, xrd_ave


# ──────────────────────────────────────────────────────────────────────────────
#  WORD TEMPLATE FILLING
# ──────────────────────────────────────────────────────────────────────────────
# Each blank is found by the visible text around it (e.g. the "____" right
# before "microns"), never by Word's internal text chunks ("runs"), which Word
# splits differently whenever the file is edited and saved. If a blank can't be
# found, generation stops with a message naming it rather than fill the wrong
# place. A value replaces its blank in a run of its own, formatted like the
# blank; values marked _OnBlank are centred on the blank's width and
# underlined, so they sit on the form's line. A missing value leaves the blank
# as printed.

_W    = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_M    = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_WP   = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_V    = "{urn:schemas-microsoft-com:vml}"
_O    = "{urn:schemas-microsoft-com:office:office}"
_W14  = "{http://schemas.microsoft.com/office/word/2010/wordml}"
_WP14 = "{http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing}"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
NBSP = " "

# The IBAD#3 box's equation labels as they appear in a paragraph's visible
# text (see _para_segments): IBAD#3 with subscript EVAL is ⟦IBAD#3EVAL⟧, etc.
_EVAL_LABEL, _REF_LABEL = "⟦IBAD#3EVAL⟧", "⟦IBAD#3ref⟧"
_MIN_LABEL,  _AVE_LABEL = "⟦XRDTiltMIN⟧", "⟦XRDTiltAVE⟧"

# Where "XRD Tilt" starts in the IBAD#3 box, as on the form's first line
# (105 pt from the box's text edge). A tab stop there keeps both XRD labels
# aligned whatever the length of the EVAL / ref values before them.
XRD_TAB_STOP = Pt(105)


class TemplateMismatch(ValueError):
    """The Word template has no blank where the program expects one."""

    def __init__(self, what):
        super().__init__(
            f"The Word template doesn't match the form this program fills: "
            f"couldn't find {what}.\n\n"
            f"Use the original {TEMPLATE_FILE}, or update the program if the "
            f"form has been revised.")


def _para_segments(p):
    """
    The visible text of paragraph element `p` as (text, w:t element) pieces.
    Only text in the paragraph's own runs is editable (it has a w:t element);
    equation labels appear as ⟦label⟧, tabs as "\\t" and line breaks as "\\n"
    (with None). Text boxes anchored in `p` have paragraphs of their own.
    """
    segs = []
    for child in p:
        if child.tag == _W + "r":
            for el in child:
                if el.tag == _W + "t":
                    segs.append((el.text or "", el))
                elif el.tag == _W + "tab":
                    segs.append(("\t", None))
                elif el.tag in (_W + "br", _W + "cr"):
                    segs.append(("\n", None))
        elif child.tag == _M + "oMath":
            label = "".join(t.text or "" for t in child.iter(_M + "t"))
            segs.append(("⟦" + "".join(label.split()) + "⟧", None))
    return segs


def _para_text(p):
    return "".join(text for text, _ in _para_segments(p))


def _find_paragraph(paragraphs, pattern, what):
    """The one paragraph element among `paragraphs` whose text matches `pattern`."""
    found = [p for p in paragraphs if re.search(pattern, _para_text(p))]
    if len(found) != 1:
        raise TemplateMismatch(what)
    return found[0]


def _value_run(like_r, value, underline=False):
    """A new run holding `value` ("\\t" = tab, "\\n" = line break), formatted
    like the run `like_r`, underlined or not as `underline` says."""
    r = OxmlElement("w:r")
    rpr = like_r.find(_W + "rPr")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    for part in re.split(r"([\t\n])", value):
        if part == "\t":
            r.append(OxmlElement("w:tab"))
        elif part == "\n":
            r.append(OxmlElement("w:br"))
        elif part:
            t = OxmlElement("w:t")
            t.text = part
            t.set(_XML_SPACE, "preserve")
            r.append(t)
    Run(r, None).font.underline = bool(underline)
    if underline:                       # pad spaces in the run's own (Latin) font
        fonts = r.find(f"{_W}rPr/{_W}rFonts")
        if fonts is not None:
            fonts.attrib.pop(_W + "hint", None)
    return r


def _drop_if_empty(r):
    """Remove run `r` if it holds no content (only formatting or empty text)."""
    if r.getparent() is None:
        return
    if all(el.tag == _W + "rPr" or (el.tag == _W + "t" and not el.text) for el in r):
        r.getparent().remove(r)


def _replace(p, start, end, value, underline=False):
    """Replace the visible text [start, end) of paragraph `p` with `value`."""
    pieces, pos = [], 0
    for text, t in _para_segments(p):
        lo, hi = max(start, pos), min(end, pos + len(text))
        if lo < hi:
            if t is None:              # a pattern group never spans labels or tabs
                raise TemplateMismatch("an editable blank")
            pieces.append((t, lo - pos, hi - pos))
        pos += len(text)
    for t, lo, hi in pieces:           # take the blank's text out
        text = t.text or ""
        t.text = text[:lo] + text[hi:]
        t.set(_XML_SPACE, "preserve")

    # Split the first run where the blank began; the value goes in between.
    first_t, cut, _ = pieces[0]
    first_r = first_t.getparent()
    rest = OxmlElement("w:r")
    rpr = first_r.find(_W + "rPr")
    if rpr is not None:
        rest.append(copy.deepcopy(rpr))
    after = first_t.text[cut:]
    first_t.text = first_t.text[:cut]
    if after:
        t = OxmlElement("w:t")
        t.text = after
        t.set(_XML_SPACE, "preserve")
        rest.append(t)
    for el in list(first_t.itersiblings()):
        rest.append(el)
    first_r.addnext(rest)
    first_r.addnext(_value_run(first_r, value, underline))
    for r in {first_r, rest, *(t.getparent() for t, _, _ in pieces)}:
        _drop_if_empty(r)


# Character widths of Calibri (and its metric twin Carlito), in 1/1000 of the
# font size; regular and italic are the same for the characters values use.
_CHAR_WIDTHS = dict(zip(
    "0123456789._- \u00A0ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz",
    [507] * 10 + [252, 498, 306, 226, 226,
     579, 544, 533, 615, 488, 459, 631, 623, 252, 319, 520, 420, 855, 646, 662,
     517, 673, 543, 459, 487, 642, 567, 890, 519, 487, 468,
     479, 525, 423, 525, 498, 305, 471, 525, 229, 239, 455, 229, 799, 525, 527,
     525, 525, 349, 391, 335, 525, 452, 715, 433, 453, 395]))


def _text_width(text):
    return sum(_CHAR_WIDTHS.get(ch, 500) for ch in text)


class _OnBlank(str):
    """A value to write centred on its blank's line, underlined."""


def _centered(value, blank):
    """`value` padded with non-breaking spaces to the width of `blank`, centred
    (a value wider than its blank is written as is)."""
    pad = max(0, round((_text_width(blank) - _text_width(value)) / _CHAR_WIDTHS[NBSP]))
    return NBSP * (pad // 2) + value + NBSP * (pad - pad // 2)


def _fill(p, pattern, values, what):
    """
    Find `pattern` exactly once in paragraph `p` and replace its groups with
    `values`, in order; a value of None leaves that group as printed.
    """
    found = list(re.finditer(pattern, _para_text(p)))
    if len(found) != 1:
        raise TemplateMismatch(what)
    m = found[0]
    for group in range(len(values), 0, -1):    # last first: earlier offsets stay put
        value = values[group - 1]
        if value is None:
            continue
        if isinstance(value, _OnBlank):
            _replace(p, *m.span(group), _centered(value, m.group(group)), underline=True)
        else:
            _replace(p, *m.span(group), value)


def _fill_buffer_line(doc, rec, continuation):
    """Hastelloy thickness, tape number and coordinates on the Buffer Tape line."""
    p = _find_paragraph(doc.element.body.iterchildren(_W + "p"), r"^Buffer Tape",
                        "the Buffer Tape line")
    # Coordinates: each number is centred on its own part of the form's
    # underlined "(___ ___-  ____      )"; "(", "-" and ")" stay put. Extra pages
    # of a long tape keep the form's blank, to be filled in by hand.
    has_coords = (not continuation and rec.get("low") is not None
                  and rec.get("high") is not None)
    _fill(p, r"\(LMO, (_+) microns.*?#\s*(_+) \((_+\s*_+)-(\s*_+\s*)\)",
          [_OnBlank(rec["hastelloy"]) if rec.get("hastelloy") else None,
           _OnBlank(rec["tape_num"]) if rec.get("tape_num") else None,
           _OnBlank(str(rec["low"])) if has_coords else None,
           _OnBlank(str(rec["high"])) if has_coords else None],
          "the thickness / tape number / coordinate blanks on the Buffer Tape line")


def _fill_ibad_box(doc, rec):
    """
    EVAL / ref values (on their underlined blank, before the form's "A") and
    XRD Tilt values (on the "______", before the form's "°") in the IBAD#3
    box. Word keeps
    two copies of a text box (for current and for older Word versions); both
    are filled. The box's two lines get a real line break, so a short value
    can't pull the second line up. On a line whose EVAL / ref value is filled
    in, a tab moves "XRD Tilt" to XRD_TAB_STOP; a line whose value is missing
    keeps the form's own blank and spacing, which only just fit before that
    position.
    """
    paras = [p for p in doc.element.body.iter(_W + "p") if _EVAL_LABEL in _para_text(p)]
    if not paras:
        raise TemplateMismatch("the IBAD#3 / XRD Tilt box")
    ev, ref = rec.get("eval_val"), rec.get("ref_ic")
    x_min, x_ave = rec.get("xrd_min"), rec.get("xrd_ave")
    pattern = (re.escape(_EVAL_LABEL) + r":\xa0([\xa0 ]+)A( +)"
               + re.escape(_MIN_LABEL) + r":(_+)°( +)"
               + re.escape(_REF_LABEL) + r":  \xa0([\xa0 ]+)A( +)"
               + re.escape(_AVE_LABEL) + r":(_+)°")
    values = [
        _OnBlank(ev) if ev else None,           # on the underlined blank, before the form's "A"
        "\t" if ev else None,
        _OnBlank(x_min) if x_min else None,     # on the "______", before the form's "°"
        "\n",
        _OnBlank(ref) if ref else None,
        "\t" if ref else None,
        _OnBlank(x_ave) if x_ave else None,
    ]
    for p in paras:
        _fill(p, pattern, values, "the IBAD#3 / XRD Tilt blanks")
        Paragraph(p, None).paragraph_format.tab_stops.add_tab_stop(XRD_TAB_STOP)


def _fill_deposition_row(doc, rec, pld_by):
    """IBAD deposition date and "Sent to PLD by:" in the IBAD Deposition row."""
    for table in doc.tables:
        resp_col = None
        for row in table.rows:
            texts = [cell.text.strip() for cell in row.cells]
            if resp_col is None:
                resp_col = next((i for i, t in enumerate(texts)
                                 if t.startswith("Responsible")), None)
            if "IBAD Deposition Date" not in texts:
                continue
            if rec.get("dep_year") and rec.get("dep_month") and rec.get("dep_day"):
                p = _find_paragraph((par._p for par in row.cells[0].paragraphs),
                                    r"202_/\s*月/\s*日", "the IBAD deposition date blank")
                _fill(p, r"(202_)/( +)月/([　 ]+)日",
                      [rec["dep_year"], " " + rec["dep_month"], " " + rec["dep_day"]],
                      "the IBAD deposition date blank")
            if pld_by:
                if resp_col is None or row.cells[resp_col].text.strip():
                    raise TemplateMismatch(
                        "an empty Responsible Person cell in the IBAD Deposition row")
                run = row.cells[resp_col].paragraphs[0].add_run(pld_by)
                run.font.size = Pt(9)                # the row's text size
            return
    raise TemplateMismatch("the IBAD Deposition Date row")


def fill_page(doc, rec, pld_by="", continuation=False):
    """
    Fill one copy of the form (a python-docx Document) for tape record `rec`.
    `continuation=True` is for a long tape's extra pages: every value is the
    same as on its first page, but the coordinates keep the form's blank.
    """
    _fill_buffer_line(doc, rec, continuation)
    _fill_ibad_box(doc, rec)
    _fill_deposition_row(doc, rec, pld_by)


def _first_free_id(doc):
    """First number above every drawing ID (wp:docPr) in `doc`."""
    return 1 + max((int(d.get("id")) for d in doc.element.body.iter(_WP + "docPr")
                    if (d.get("id") or "").isdigit()), default=0)


def _append_page(out_doc, page_doc, next_id):
    """
    Append the filled form `page_doc` to `out_doc` as a new page. The last
    paragraph of `out_doc` becomes a section break (an extra page-break
    paragraph would push a blank page), and the IDs that must be unique within
    one Word file are renumbered from `next_id` on; the text/paragraph IDs are
    dropped and Word regenerates them. Returns the next free ID number.
    """
    body = out_doc.element.body
    sect = body.find(_W + "sectPr")
    last_p = sect.getprevious()
    if last_p is None or last_p.tag != _W + "p":
        raise TemplateMismatch("the empty paragraph that ends the form")
    ppr = last_p.find(_W + "pPr")
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        last_p.insert(0, ppr)
    ppr.append(copy.deepcopy(sect))

    for child in list(page_doc.element.body):
        if child.tag == _W + "sectPr":
            continue
        for el in child.iter():
            if el.tag == _WP + "docPr":
                el.set("id", str(next_id))
                next_id += 1
            elif el.tag.startswith(_V) and el.get("id"):
                el.set("id", f"{el.get('id')}_{next_id}")
                next_id += 1
            if el.get(_O + "spid"):
                el.set(_O + "spid", f"_x0000_s{next_id}")
                next_id += 1
            for attr in (_W14 + "paraId", _W14 + "textId", _WP14 + "anchorId"):
                el.attrib.pop(attr, None)
        sect.addprevious(child)
    return next_id


# ──────────────────────────────────────────────────────────────────────────────
#  CORE GENERATION
# ──────────────────────────────────────────────────────────────────────────────

def read_monday(excel_path: str, log_cb=None):
    """
    Monday.com / IBAD database export reader.
    Returns (records, skipped) where each record is a normalized per-tape dict.
    Behaviour is identical to the original program's reading logic, except that
    the page count now follows the "Length, m" column, as in the RR124 reader.
    """
    header_row = find_header_row(excel_path)
    df = pd.read_excel(excel_path, engine="openpyxl", header=header_row)

    name_col       = _col(df, "Name")
    hastelloy_col  = _col(df, "Hastelloy thickness")
    eval_start_col = _col(df, "Evaluation_start")
    eval_end_col   = _col(df, "Evaluation_end")
    ref_col        = _col(df, "Reference Ic")
    dep_date_col   = _col(df, "IBAD deposition date")
    len_col        = _length_col(df)              # optional – drives page count

    missing = [n for n, c_ in [
        ("Name", name_col),
        ("Hastelloy thickness", hastelloy_col),
        ("Reference Ic", ref_col),
    ] if c_ is None]
    if missing:
        raise ValueError(
            "The following column(s) were not found in the Excel file:\n\n"
            + "\n".join(f"  • {m}" for m in missing)
            + "\n\nPlease ensure the file contains these exact column names."
        )

    if eval_start_col is None and eval_end_col is None and log_cb:
        log_cb("⚠  Neither Evaluation_start nor Evaluation_end found — "
               "EVAL field will be blank.")
    if dep_date_col is None and log_cb:
        log_cb("⚠  'IBAD deposition date' column not found — "
               "deposition date will be blank.")
    if len_col is None and log_cb:
        log_cb("⚠  'Length, m' column not found — every tape will get 1 page.")

    records, skipped = [], []
    for i, (_, row) in enumerate(df.iterrows()):
        raw_name = str(row[name_col]) if name_col else ""

        if not re.search(r"#[A-Za-z0-9]+\s*\(", raw_name):
            if raw_name.strip() and raw_name.lower() != "nan":
                skipped.append(
                    f"Row {i + header_row + 2}: no tape pattern  –  \"{raw_name[:55]}\"")
            continue

        tape_num, low, high = parse_name(raw_name)
        if tape_num is None:
            skipped.append(
                f"Row {i + header_row + 2}: could not parse coords  –  \"{raw_name[:55]}\"")
            continue

        hastelloy = fmt_num(row[hastelloy_col]) if hastelloy_col else ""
        eval_val  = pick_eval(
            row[eval_start_col] if eval_start_col else None,
            row[eval_end_col]   if eval_end_col   else None,
        )
        ref_ic = fmt_num(row[ref_col]) if ref_col else ""

        dep_year, dep_month, dep_day = (None, None, None)
        if dep_date_col is not None:
            dep_year, dep_month, dep_day = fmt_deposition_date(row[dep_date_col])

        length, n_pages = _length_and_pages(row, len_col)

        dep_str = f"{dep_year}/{dep_month}/{dep_day}" if dep_year else "—"
        len_str = f"{fmt_float(length)} m" if length is not None else "—"
        records.append(dict(
            tape_num=tape_num, low=low, high=high, hastelloy=hastelloy,
            eval_val=eval_val, ref_ic=ref_ic,
            dep_year=dep_year, dep_month=dep_month, dep_day=dep_day,
            xrd_min="", xrd_ave="",
            length=length, pages=n_pages,
            log=(f"✔  #{tape_num} ({low}–{high})   "
                 f"Hastelloy={hastelloy} µm   "
                 f"EVAL={eval_val or '—'} A   "
                 f"REF={ref_ic or '—'} A   "
                 f"Date={dep_str}   "
                 f"Len={len_str} → {n_pages} pg"),
        ))

    return records, skipped


def read_rr124(excel_path: str, selected_date, log_cb=None):
    """
    RR124-format reader. One record per 'Sent to PLD' row whose Cut_Date matches
    `selected_date`. IC / XRD / Reference_Ic are pulled from the matching
    cut-eval samples per the finalized feature spec.
    Returns (records, skipped).
    """
    if not selected_date:
        raise ValueError("A Cut_Date must be selected before generating "
                         "from an RR124-format file.")

    hdr = _find_rr124_header(excel_path)
    df  = pd.read_excel(excel_path, engine="openpyxl", header=hdr)

    run_col    = _col(df, "Run_ID")
    name_col   = _col(df, "Tape_name")
    dep_col    = _col(df, "Date_of_IBAD_Deposition")
    thick_col  = _col(df, "Hastealloy_Thickness")   # note: misspelled in export
    status_col = _col(df, "Tape_Status")
    cut_col    = _col(df, "Cut_Date")
    s_col      = _col(df, "Sample_Cuts_Start_Coordinate")
    e_col      = _col(df, "Sample_Cuts_End_Coordinate")
    ic_col     = _col(df, "Ic")
    ref_col    = _col(df, "Reference_Ic")
    len_col    = _length_col(df)                    # optional – drives page count
    xrd_cols   = [c for c in df.columns
                  if str(c).strip().lower().startswith("2d_xrd_tilt_")]

    required = {
        "Run_ID": run_col, "Tape_name": name_col, "Tape_Status": status_col,
        "Cut_Date": cut_col, "Sample_Cuts_Start_Coordinate": s_col,
        "Sample_Cuts_End_Coordinate": e_col, "Ic": ic_col,
        "Reference_Ic": ref_col, "Date_of_IBAD_Deposition": dep_col,
        "Hastealloy_Thickness": thick_col,
    }
    missing = [k for k, v in required.items() if v is None]
    if missing:
        raise ValueError(
            "The RR124 file is missing these column(s):\n\n"
            + "\n".join(f"  • {m}" for m in missing))
    if len(xrd_cols) < 3 and log_cb:
        log_cb(f"⚠  Found {len(xrd_cols)} XRD tilt column(s); expected 3 — "
               "XRD fields may stay blank.")
    if len_col is None and log_cb:
        log_cb("⚠  'Length' column not found — every tape will get 1 page.")

    rows = df.to_dict("records")

    # Candidate pools keyed on (Run_ID, base tape name) over cut-eval rows.
    pool = defaultdict(list)
    for r in rows:
        if _status_norm(r[status_col]) == "cut eval sample":
            pool[(r[run_col], _base_name(r[name_col]))].append(r)

    sel = pd.to_datetime(selected_date, errors="coerce")
    sel = sel.date() if not pd.isna(sel) else None

    records, skipped = [], []
    for r in rows:
        if _status_norm(r[status_col]) != "sent to pld":
            continue
        cd = pd.to_datetime(r[cut_col], errors="coerce")
        if pd.isna(cd) or sel is None or cd.date() != sel:
            continue

        base = _base_name(r[name_col])
        s = rr_to_float(r[s_col])
        e = rr_to_float(r[e_col])
        if s is not None and e is not None:
            low, high = int(round(min(s, e))), int(round(max(s, e)))
        elif s is not None:
            low = high = int(round(s))
        elif e is not None:
            low = high = int(round(e))
        else:
            low = high = None

        hastelloy = fmt_num(r[thick_col])
        dep_year, dep_month, dep_day = fmt_deposition_date(r[dep_col])

        key   = (r[run_col], base)
        cands = _select_candidates(pool[key], s, e, s_col, e_col, ic_col)
        eval_val, ref_ic, xrd_min, xrd_ave = _derive_values(
            cands, pool[key], s, e, s_col, e_col, ic_col, ref_col, xrd_cols)

        length, n_pages = _length_and_pages(r, len_col)

        dep_str = f"{dep_year}/{dep_month}/{dep_day}" if dep_year else "—"
        len_str = f"{fmt_float(length)} m" if length is not None else "—"
        records.append(dict(
            tape_num=base, low=low, high=high, hastelloy=hastelloy,
            eval_val=eval_val, ref_ic=ref_ic,
            dep_year=dep_year, dep_month=dep_month, dep_day=dep_day,
            xrd_min=xrd_min, xrd_ave=xrd_ave,
            length=length, pages=n_pages,
            log=(f"✔  {base} ({low}–{high})   "
                 f"Hastelloy={hastelloy} µm   "
                 f"EVAL={eval_val or '—'} A   "
                 f"REF={ref_ic or '—'} A   "
                 f"XRD_MIN={xrd_min or '—'}  XRD_AVE={xrd_ave or '—'}  "
                 f"Date={dep_str}   "
                 f"Len={len_str} → {n_pages} pg"),
        ))

    return records, skipped


def generate(excel_path: str, template_path: str, output_path: str,
             pld_by: str = "", selected_date=None, fmt: str = None,
             progress_cb=None, log_cb=None):
    """
    Detect the file format (unless `fmt` is supplied), dispatch to the right
    reader, then fill the Word template once for every page each record calls
    for (see pages_for_length) and save all pages as one Word file. Page 1
    carries the coordinates; any further page for the same tape repeats every
    other value with the coordinate field left blank.

    Returns (pages_generated, skipped_descriptions).
    """
    if not template_path.lower().endswith(".docx"):
        raise ValueError(
            "The template must be the Word form (.docx):\n"
            f"{TEMPLATE_FILE}")
    if fmt is None:
        fmt = detect_format(excel_path)

    if fmt == "rr124":
        records, skipped = read_rr124(excel_path, selected_date, log_cb=log_cb)
    elif fmt == "monday":
        records, skipped = read_monday(excel_path, log_cb=log_cb)
    else:
        raise ValueError(f"Unknown format: {fmt!r}")

    if not records:
        if fmt == "rr124":
            raise ValueError(
                "No 'Sent to PLD' rows match the selected Cut_Date.\n\n"
                "Pick a different date, or check the file.")
        raise ValueError(
            "No valid tape rows were found in the Excel file.\n\n"
            "Make sure the Name column contains entries like:\n"
            "  #3244(1235–2165)\n  #KM3268(450–10)")

    # Read the template once; every page starts from a fresh copy of it.
    with open(template_path, "rb") as fh:
        template_bytes = fh.read()
    out_doc = None
    # Progress is counted in PAGES, not records: a single long tape can now
    # contribute up to MAX_PAGES pages.
    total = sum(max(1, int(rec.get("pages", 1))) for rec in records)
    ok    = 0

    for rec in records:
        n_pages = max(1, min(MAX_PAGES, int(rec.get("pages", 1))))

        if log_cb and rec.get("log"):
            log_cb(rec["log"])

        for p in range(n_pages):
            first = (p == 0)

            page = docx.Document(io.BytesIO(template_bytes))
            fill_page(page, rec, pld_by, continuation=not first)
            if out_doc is None:
                out_doc = page
                next_id = _first_free_id(out_doc)
            else:
                next_id = _append_page(out_doc, page, next_id)
            ok += 1

            if progress_cb:
                progress_cb(ok, total)
            if log_cb and n_pages > 1 and not first:
                log_cb(f"   ↳ page {p + 1} of {n_pages} — coordinates blank")

    try:
        out_doc.save(output_path)
    except PermissionError:
        raise PermissionError(
            f"Couldn't save {output_path}\n\n"
            "If that file is open in Word, close it and try again.") from None

    if log_cb:
        log_cb(f"\n{len(records)} tape(s) → {ok} page(s).")

    return ok, skipped


# ──────────────────────────────────────────────────────────────────────────────
#  GUI
# ──────────────────────────────────────────────────────────────────────────────

ACCENT  = "#803C43"   # light maroon
BG      = "#F5F7FA"
BTN_BG  = "#803C43"   # light maroon
BTN_HVR = "#5E2C31"   # darker maroon for hover
BTN_FG  = "#FFFFFF"
GRAY    = "#90A4AE"
TEXT_DK = "#212121"
TEXT_SM = "#546E7A"
GREEN   = "#2E7D32"
RED     = "#C62828"


class App(tk.Tk):

    def __init__(self):
        super().__init__()
        self.title("Tape Process Form Generator")
        self.resizable(False, False)
        self.configure(bg=BG)
        self._set_icon()

        self.excel_path    = tk.StringVar()
        self.template_path = tk.StringVar()
        self.output_path   = tk.StringVar()
        self.pld_by        = tk.StringVar()
        self.cut_date      = tk.StringVar()

        self.detected_format = None     # 'rr124' | 'monday' | None
        self._cut_dates      = []       # list[date] available for the picker

        default = os.path.join(_app_dir(), TEMPLATE_FILE)
        if os.path.isfile(default):
            self.template_path.set(default)

        self._build_ui()
        self._center_window(510, 648)

    def _set_icon(self):
        """Show the program logo in the title bar and on the taskbar."""
        icon = _bundled_file(ICON_FILE)
        if not os.path.isfile(icon):
            return
        try:
            self.iconbitmap(default=icon)
        except tk.TclError:
            pass        # Tk reads .ico files only on Windows

    # ── Input validation ────────────────────────────────────────────────────
    def _validate_pld(self, proposed: str) -> bool:
        """Allow typing but cap length at PLD_MAX_CHARS (min is checked on Generate)."""
        return len(proposed) <= PLD_MAX_CHARS

    # ── Layout ────────────────────────────────────────────────────────────────

    def _build_ui(self):
        hdr = tk.Frame(self, bg=ACCENT, height=56)
        hdr.pack(fill="x")
        tk.Label(
            hdr, text="IBAD Tape Process Form",
            font=("Segoe UI", 14, "bold"),
            fg="white", bg=ACCENT, anchor="center",
        ).pack(fill="x", pady=14)

        body = tk.Frame(self, bg=BG, padx=20, pady=16)
        body.pack(fill="both", expand=True)
        body.grid_columnconfigure(1, weight=1)

        self._file_row(body, "Excel Data File (.xlsx):",
                       self.excel_path, self._pick_excel, row=0)
        self._file_row(body, "Template (IBAD-PF-009 .docx):",
                       self.template_path, self._pick_template, row=1)
        self._file_row(body, "Save Output Word File As:",
                       self.output_path, self._pick_output,
                       row=2, is_save=True)

        # 4th input row: "Sent to PLD by:" (typed value, not a file)
        tk.Label(
            body, text="Sent to PLD by:", font=("Segoe UI", 9),
            fg=TEXT_DK, bg=BG, anchor="w",
        ).grid(row=3, column=0, sticky="w", padx=(0, 8), pady=3)
        vcmd = (self.register(self._validate_pld), "%P")
        self.pld_entry = tk.Entry(
            body, textvariable=self.pld_by, width=34,
            font=("Segoe UI", 9), relief="solid", bd=1,
            validate="key", validatecommand=vcmd,
        )
        self.pld_entry.grid(row=3, column=1, sticky="ew", pady=3)
        tk.Label(
            body, text=f"{PLD_MIN_CHARS}–{PLD_MAX_CHARS} chars",
            font=("Segoe UI", 8), fg=GRAY, bg=BG, anchor="w",
        ).grid(row=3, column=2, sticky="w", padx=(6, 0), pady=3)

        # 5th input row: Cut_Date picker — shown only for RR124-format files.
        self.cut_date_label = tk.Label(
            body, text="Date sent to PLD:", font=("Segoe UI", 9),
            fg=TEXT_DK, bg=BG, anchor="w",
        )
        self.cut_date_label.grid(row=4, column=0, sticky="w", padx=(0, 8), pady=3)
        self.cut_date_combo = ttk.Combobox(
            body, textvariable=self.cut_date, width=32,
            font=("Segoe UI", 9), state="readonly",
        )
        self.cut_date_combo.grid(row=4, column=1, sticky="ew", pady=3)
        self.cut_date_hint = tk.Label(
            body, text="required", font=("Segoe UI", 8),
            fg=GRAY, bg=BG, anchor="w",
        )
        self.cut_date_hint.grid(row=4, column=2, sticky="w", padx=(6, 0), pady=3)
        # Hidden until an RR124 file is selected.
        self.cut_date_label.grid_remove()
        self.cut_date_combo.grid_remove()
        self.cut_date_hint.grid_remove()

        ttk.Separator(body, orient="horizontal").grid(
            row=6, column=0, columnspan=3, sticky="ew", pady=10)

        self.gen_btn = self._btn(body, "⚙   Generate Word File",
                                 self._start_generate, width=22)
        self.gen_btn.grid(row=7, column=0, columnspan=3, pady=(0, 6))

        self.progress = ttk.Progressbar(body, length=440, mode="determinate")
        self.progress.grid(row=8, column=0, columnspan=3,
                           sticky="ew", pady=(4, 2))

        self.status_var = tk.StringVar(value="Ready – select an Excel file to begin.")
        self._status_lbl = tk.Label(
            body, textvariable=self.status_var,
            font=("Segoe UI", 9), fg=TEXT_SM, bg=BG,
            wraplength=440, justify="left", anchor="w",
        )
        self._status_lbl.grid(row=9, column=0, columnspan=3,
                               sticky="w", pady=(2, 4))

        log_frame = tk.Frame(body, bg=BG)
        log_frame.grid(row=10, column=0, columnspan=3,
                       sticky="nsew", pady=(6, 0))
        body.grid_rowconfigure(10, weight=1)

        self.log_box = tk.Text(
            log_frame, height=9, width=60,
            font=("Consolas", 8), bg="#ECEFF1",
            relief="flat", state="disabled", wrap="word",
        )
        sb = ttk.Scrollbar(log_frame, command=self.log_box.yview)
        self.log_box.configure(yscrollcommand=sb.set)
        self.log_box.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

        tk.Label(
            self,
            text="Faraday Factory Japan",
            font=("Segoe UI", 8), fg=GRAY, bg=BG,
        ).pack(pady=(4, 8))

    def _file_row(self, parent, label, var, cmd, row, is_save=False):
        tk.Label(
            parent, text=label, font=("Segoe UI", 9),
            fg=TEXT_DK, bg=BG, anchor="w",
        ).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=3)
        tk.Entry(
            parent, textvariable=var, width=34,
            font=("Segoe UI", 9), relief="solid", bd=1,
        ).grid(row=row, column=1, sticky="ew", pady=3)
        self._btn(parent, "Save As…" if is_save else "Browse…",
                  cmd, width=10
                  ).grid(row=row, column=2, padx=(6, 0), pady=3)

    def _btn(self, parent, text, cmd, width=None):
        kw = {"width": width} if width else {}
        b = tk.Button(
            parent, text=text, command=cmd,
            font=("Segoe UI", 9, "bold"),
            bg=BTN_BG, fg=BTN_FG,
            activebackground=BTN_HVR, activeforeground=BTN_FG,
            relief="flat", cursor="hand2", padx=10, pady=5, **kw,
        )
        b.bind("<Enter>", lambda _: b.configure(bg=BTN_HVR))
        b.bind("<Leave>", lambda _: b.configure(bg=BTN_BG))
        return b

    def _center_window(self, w, h):
        self.update_idletasks()
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    # ── File pickers ──────────────────────────────────────────────────────────

    def _pick_excel(self):
        path = filedialog.askopenfilename(
            title="Select Excel Data File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if path:
            self.excel_path.set(path)
            if not self.output_path.get():
                base = os.path.splitext(path)[0]
                self.output_path.set(base + "_TapeProcessForms.docx")
            self._log(f"Excel: {os.path.basename(path)}")
            self._detect_and_setup(path)

    def _detect_and_setup(self, path):
        """Detect the file format and show/hide the Cut_Date picker."""
        self.detected_format = None
        self._cut_dates = []
        self.cut_date.set("")
        try:
            self.detected_format = detect_format(path)
        except Exception as exc:
            self._log(f"⚠  Could not detect format: {exc}")

        if self.detected_format == "rr124":
            try:
                self._cut_dates = rr124_cut_dates(path)
            except Exception as exc:
                self._log(f"⚠  Could not read Cut_Date values: {exc}")
            self._log(f"Format: RR124  –  {len(self._cut_dates)} cut date(s) available.")
        elif self.detected_format == "monday":
            self._log("Format: Monday (existing reader).")

        self._update_cut_date_picker()

    def _update_cut_date_picker(self):
        """Reveal the picker (populated with available dates) only for RR124."""
        if self.detected_format == "rr124" and self._cut_dates:
            self.cut_date_combo["values"] = [d.isoformat() for d in self._cut_dates]
            self.cut_date_label.grid()
            self.cut_date_combo.grid()
            self.cut_date_hint.grid()
        else:
            self.cut_date_label.grid_remove()
            self.cut_date_combo.grid_remove()
            self.cut_date_hint.grid_remove()

    def _pick_template(self):
        path = filedialog.askopenfilename(
            title="Select Tape Process Form Template (Word file)",
            filetypes=[("Word documents", "*.docx"), ("All files", "*.*")],
        )
        if path:
            self.template_path.set(path)
            self._log(f"Template: {os.path.basename(path)}")

    def _pick_output(self):
        init_file = (os.path.basename(self.output_path.get())
                     if self.output_path.get() else "TapeProcessForms.docx")
        init_dir  = (os.path.dirname(self.output_path.get())
                     if self.output_path.get() else os.path.expanduser("~"))
        path = filedialog.asksaveasfilename(
            title="Save Output Word File As",
            defaultextension=".docx",
            initialfile=init_file,
            initialdir=init_dir,
            filetypes=[("Word documents", "*.docx")],
        )
        if path:
            self.output_path.set(path)

    # ── Log / status ──────────────────────────────────────────────────────────

    def _log(self, msg: str):
        self.log_box.configure(state="normal")
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    def _set_status(self, msg: str, color: str = TEXT_SM):
        self.status_var.set(msg)
        self._status_lbl.configure(fg=color)

    # ── Generation ────────────────────────────────────────────────────────────

    def _start_generate(self):
        excel = self.excel_path.get().strip()
        tpl   = self.template_path.get().strip()
        out   = self.output_path.get().strip()
        pld   = self.pld_by.get().strip()
        sel_date = self.cut_date.get().strip()

        if not excel:
            messagebox.showwarning("Missing File",
                                   "Please select an Excel data file.")
            return
        if not tpl:
            messagebox.showwarning(
                "Missing Template",
                "Word template not found.\n\n"
                f"Place {TEMPLATE_FILE} in the same\n"
                "folder as this program, or use Browse… to locate it.")
            return
        if not out:
            messagebox.showwarning("Missing Path",
                                   "Please specify where to save the output Word file.")
            return
        # "Sent to PLD by:" — required, 4–10 characters
        if not (PLD_MIN_CHARS <= len(pld) <= PLD_MAX_CHARS):
            messagebox.showwarning(
                "Check \u201CSent to PLD by\u201D",
                f"The \u201CSent to PLD by:\u201D field must be between "
                f"{PLD_MIN_CHARS} and {PLD_MAX_CHARS} characters.\n\n"
                f"It currently has {len(pld)} character(s).")
            self.pld_entry.focus_set()
            return
        # RR124 files require a Cut_Date selection before generating
        if self.detected_format == "rr124" and not sel_date:
            messagebox.showwarning(
                "Select a Cut Date",
                "This is an RR124-format file, so a Cut Date is required.\n\n"
                "Choose a date from the dropdown before generating.")
            self.cut_date_combo.focus_set()
            return
        if not os.path.isfile(excel):
            messagebox.showerror("File Not Found", f"Excel file not found:\n{excel}")
            return
        if not os.path.isfile(tpl):
            messagebox.showerror("File Not Found", f"Word template not found:\n{tpl}")
            return
        # The output is a Word file: make sure its name says so. Only a real
        # document extension (e.g. an old ".pdf" name) is swapped; anything
        # else is kept whole, so "Forms 2026.10.14" becomes
        # "Forms 2026.10.14.docx", not "Forms 2026.10.docx".
        if not out.lower().endswith(".docx"):
            root, ext = os.path.splitext(out)
            if ext.lower() in (".pdf", ".doc", ".docm", ".dot", ".dotx", ".dotm"):
                out = root
            out += ".docx"
            self.output_path.set(out)

        self.gen_btn.configure(state="disabled", text="⏳  Generating…")
        self.progress["value"] = 0
        self._log(f"\n{'─' * 52}")
        self._log("Starting…")
        self._set_status("Generating the Word file, please wait…")

        fmt = self.detected_format

        def _worker():
            try:
                def _prog(done, total):
                    self.progress["value"] = int(done / total * 100)
                    self._set_status(f"Processing page {done} of {total}…")

                pages, skipped = generate(excel, tpl, out,
                                          pld_by=pld,
                                          selected_date=sel_date or None,
                                          fmt=fmt,
                                          progress_cb=_prog,
                                          log_cb=self._log)
                self.progress["value"] = 100

                if skipped:
                    self._log(f"\n⚠  {len(skipped)} row(s) skipped:")
                    for s in skipped:
                        self._log(f"   {s}")

                self._log(f"\n✅  {pages} page(s) saved to:\n   {out}")
                self._set_status(
                    f"✅  Done!  {pages} page(s) →  {os.path.basename(out)}",
                    color=GREEN,
                )

                if messagebox.askyesno(
                    "Word File Generated",
                    f"✅  Success!\n\n"
                    f"{pages} page(s) saved to:\n{out}\n\n"
                    "Open the output folder?",
                    icon="info",
                ):
                    _open_folder(os.path.dirname(os.path.abspath(out)))

            except Exception as exc:
                self._log(f"\n❌  Error: {exc}")
                self._set_status(f"Error: {exc}", color=RED)
                messagebox.showerror("Generation Failed", str(exc))
            finally:
                self.gen_btn.configure(state="normal", text="⚙   Generate Word File")

        threading.Thread(target=_worker, daemon=True).start()


def _open_folder(path: str):
    import platform, subprocess
    sys_ = platform.system()
    if sys_ == "Windows":
        os.startfile(path)
    elif sys_ == "Darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


# ──────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    # Run as a .pyw, the taskbar would show pythonw.exe's icon; an app ID of
    # its own makes it use the window's logo. Not in the built .exe: that
    # already carries the logo, and an explicit ID would stop it grouping
    # with a taskbar pin of the .exe (a second button would appear).
    if not getattr(sys, "frozen", False):
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                "FaradayFactoryJapan.TapeProcessFormGenerator")
        except Exception:
            pass    # not Windows
    app = App()
    app.mainloop()
